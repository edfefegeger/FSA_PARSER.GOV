#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Парсер реестра ФСА (pub.fsa.gov.ru): сертификаты и декларации о соответствии.

Выбирает документы, срок действия которых истекает в ближайшие N дней,
фильтрует по номеру технического регламента (ТР ТС / ТР ЕАЭС / ТР РФ)
и выгружает в Excel / CSV.

Запуск:
    python fsa_parser.py                 # по настройкам из config.ini
    python fsa_parser.py --days 90       # разово изменить порог в днях
    python fsa_parser.py --tr-ts 017,021 # разово выбрать другие техрегламенты
    python fsa_parser.py --list-trts     # показать все номера техрегламентов
    python fsa_parser.py --debug         # сохранить «сырые» ответы сайта для диагностики
"""
import argparse
import configparser
import csv
import json
import os
import random
import re
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import requests

# ----------------------------------------------------------------------------
# Консоль: UTF-8 и цвета (в т.ч. в Windows 10/11)
# ----------------------------------------------------------------------------
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass
if os.name == "nt":
    os.system("")  # включает ANSI-последовательности в cmd

C = {"b": "\033[94m", "g": "\033[92m", "y": "\033[93m", "r": "\033[91m",
     "d": "\033[90m", "x": "\033[0m"}


def paint(text, color):
    return f"{C[color]}{text}{C['x']}"


# ----------------------------------------------------------------------------
# Константы
# ----------------------------------------------------------------------------
VERSION = "1.1"          # 1.1 - фильтр по номеру ТР ТС и столбец «Наименование ТР ТС»
BASE = "https://pub.fsa.gov.ru"
ANON_LOGIN = {"username": "anonymous", "password": "hrgesf7HDR67Bd"}  # публичный гостевой вход сайта

DOC_KINDS = {
    "certificates": {
        "title": "Сертификат",
        "list_url": f"{BASE}/api/v1/rss/common/certificates/get",
        "detail_url": f"{BASE}/api/v1/rss/common/certificates/{{id}}",
        "sort_column": "endDate",
    },
    "declarations": {
        "title": "Декларация",
        "list_url": f"{BASE}/api/v1/rds/common/declarations/get",
        "detail_url": f"{BASE}/api/v1/rds/common/declarations/{{id}}",
        "sort_column": "declEndDate",  # у сущности SearchDeclaration нет поля endDate
    },
}
STATUS_ACTIVE = 6  # статус «Действует»

# Технические регламенты: то же, что в фильтре «Технический регламент» на сайте.
# Значение (число) - внутренний id реестра, он уходит в запрос idTechReg.
TR_TS_IDS = {
    "ТР ТС 001/2011": 33,
    "ТР ТС 002/2011": 18,
    "ТР ТС 003/2011": 37,
    "ТР ТС 004/2011": 11,
    "ТР ТС 005/2011": 38,
    "ТР ТС 006/2011": 1,
    "ТР ТС 007/2011": 39,
    "ТР ТС 008/2011": 17,
    "ТР ТС 009/2011": 28,
    "ТР ТС 010/2011": 14,
    "ТР ТС 011/2011": 6,
    "ТР ТС 012/2011": 12,
    "ТР ТС 013/2011": 31,
    "ТР ТС 014/2011": 29,
    "ТР ТС 015/2011": 32,
    "ТР ТС 016/2011": 16,
    "ТР ТС 017/2011": 8,
    "ТР ТС 018/2011": 30,
    "ТР ТС 019/2011": 36,
    "ТР ТС 020/2011": 24,
    "ТР ТС 021/2011": 27,
    "ТР ТС 022/2011": 97161,
    "ТР ТС 023/2011": 22,
    "ТР ТС 024/2011": 20,
    "ТР ТС 025/2012": 4521250,
    "ТР ТС 026/2012": 40,
    "ТР ТС 027/2012": 97185,
    "ТР ТС 028/2012": 26,
    "ТР ТС 029/2012": 23,
    "ТР ТС 030/2012": 34,
    "ТР ТС 031/2012": 10,
    "ТР ТС 032/2013": 5,
    "ТР ТС 033/2013": 21,
    "ТР ТС 034/2013": 2,
    "ТР ТС 035/2014": 35,
    "ТР ЕАЭС 036/2016": 97149,
    "ТР ЕАЭС 037/2016": 4501729,
    "ТР ЕАЭС 038/2016": 97167,
    "ТР ЕАЭС 039/2016": 97156,
    "ТР ЕАЭС 040/2016": 15,
    "ТР ЕАЭС 041/2017": 97209,
    "ТР ЕАЭС 042/2017": 97210,
    "ТР ЕАЭС 043/2017": 97211,
    "ТР ЕАЭС 044/2017": 97208,
    "ТР ЕАЭС 045/2017": 99333,
    "ТР ЕАЭС 046/2018": 4507056,
    "ТР ЕАЭС 047/2018": 101330,
    "ТР ЕАЭС 048/2019": 4508584,
    "ТР ЕАЭС 050/2021": 4517710,
    "ТР ЕАЭС 051/2021": 4516872,
    "ТР ЕАЭС 052/2021": 4522924,
    "ТР РФ 007/2008": 49,
    "ТР РФ 015/2009": 97187,
    "ТР РФ 016/2010": 97170,
    "ТР РФ 022/2010": 97163,
    "ТР РФ 023/2010": 60,
    "ТР РФ 024/2010": 97169,
    "123-ФЗ": 56,
    "126-ФЗ": 98706,
    "Решение Совета ЕЭК от 18.04.2018 № 44": 4511534,
    "ГОСТ Р 54008-2022": 4518156,
}
TR_TS_BY_ID = {v: k for k, v in TR_TS_IDS.items()}
TR_TS_NUM_RE = re.compile(r"(\d{3})\s*/\s*(\d{4})")
TR_TS_IN_TEXT_RE = re.compile(r"ТР\s*(?:ТС|ЕАЭС|РФ)?\s*\d{3}\s*/\s*\d{4}|\b\d{3}\s*/\s*\d{4}\b", re.I)


def tr_ts_parts(label):
    """('017', '2011') или None."""
    m = TR_TS_NUM_RE.search(label)
    return (m[1], m[2]) if m else None


def normalize_tr_ts(raw):
    """'017', '017/2011', 'ТР ТС 017/2011' -> 'ТР ТС 017/2011'. None - если такого нет."""
    s = " ".join(str(raw or "").split())
    if not s:
        return None
    flat = re.sub(r"[\s\u00a0]", "", s).upper()
    for label in TR_TS_IDS:                                  # точное совпадение со справочником
        if re.sub(r"[\s\u00a0]", "", label).upper() == flat:
            return label
    m = TR_TS_NUM_RE.search(s)
    if m:                                                    # указан номер с годом: 017/2011
        exact = [lbl for lbl in TR_TS_IDS if tr_ts_parts(lbl) == (m[1], m[2])]
        return exact[0] if len(exact) == 1 else None
    m = re.fullmatch(r"(?:ТР[\s-]*(?:ТС|ЕАЭС|РФ)?[\s-]*)?(\d{3})", flat)
    if m:                                                    # указан только номер: 017
        same = [lbl for lbl in TR_TS_IDS if tr_ts_parts(lbl) and tr_ts_parts(lbl)[0] == m[1]]
        return same[0] if len(same) == 1 else None
    return None


def tr_ts_hint():
    """Подсказка со всеми номерами - когда в config.ini указан неизвестный регламент."""
    return ", ".join(lbl for lbl in TR_TS_IDS)


def extract_tr_ts(det):
    """Технические регламенты из карточки документа (если сайт их там отдаёт)."""
    found = []

    def walk(obj, in_trts=False):
        if isinstance(obj, dict):
            for k, v in obj.items():
                hint = in_trts or bool(re.search(r"techreg|trts|tr_ts|reglament|техрегл", str(k), re.I))
                walk(v, hint)
        elif isinstance(obj, list):
            for v in obj:
                walk(v, in_trts)
        elif in_trts and isinstance(obj, int) and not isinstance(obj, bool):
            if obj in TR_TS_BY_ID:
                found.append(TR_TS_BY_ID[obj])
        elif in_trts and isinstance(obj, str):
            for m in TR_TS_IN_TEXT_RE.finditer(obj):
                found.append(normalize_tr_ts(m.group(0)) or m.group(0).strip())

    walk(det)
    return "; ".join(sorted(set(found)))


COLUMN_TITLES = {
    "company": "Наименование юрлица",
    "head": "ФИО руководителя",
    "last_name": "Фамилия",
    "first_name": "Имя",
    "patronymic": "Отчество",
    "number": "Номер документа",
    "doc_type": "Тип документа",
    "end_date": "Дата окончания",
    "days_left": "Дней до конца",
    "email": "Электронный адрес",
    "tr_ts": "Наименование ТР ТС",
    "product": "Наименование продукции",
}
COLUMN_WIDTHS = {"company": 34, "head": 26, "last_name": 18, "first_name": 16, "patronymic": 20, "number": 34, "doc_type": 14,
                 "end_date": 16, "days_left": 14, "email": 30, "tr_ts": 22, "product": 50}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")


# ----------------------------------------------------------------------------
# Настройки
# ----------------------------------------------------------------------------
def load_config(path):
    cp = configparser.ConfigParser(inline_comment_prefixes=None)
    if not Path(path).exists():
        sys.exit(f"Не найден файл настроек: {path}")
    cp.read(path, encoding="utf-8-sig")

    def lst(section, key, default=""):
        raw = cp.get(section, key, fallback=default)
        return [x.strip() for x in raw.split(",") if x.strip()]

    def tr_list():
        """Техрегламенты из config.ini -> обозначения вида «ТР ТС 017/2011»."""
        out = []
        for raw in lst("filter", "tr_ts"):
            label = normalize_tr_ts(raw)
            if not label:
                sys.exit(f"Не понял технический регламент «{raw}» в config.ini (tr_ts).\n"
                         f"  Пишите номер: 017 или 017/2011 (несколько - через запятую).\n"
                         f"  Доступные номера: {tr_ts_hint()}")
            if label not in out:
                out.append(label)
        return out

    cfg = {
        "days": cp.getint("filter", "days_threshold", fallback=60),
        "tr_ts": tr_list(),
        "keywords": [k.lower() for k in lst("filter", "product_keywords")],
        "kinds": [k for k in lst("filter", "document_types", "certificates,declarations")
                  if k in DOC_KINDS],
        "format": cp.get("output", "format", fallback="xlsx").strip().lower(),
        "out_dir": cp.get("output", "output_dir", fallback="выгрузки").strip(),
        "prefix": cp.get("output", "file_prefix", fallback="fsa").strip(),
        "columns": [c for c in lst("output", "columns") if c in COLUMN_TITLES]
                   or list(COLUMN_TITLES),
        "tr_ts_doc": cp.getboolean("output", "tr_ts_from_document", fallback=False),
        "name_case": cp.get("output", "name_case", fallback="title").strip().lower(),
        "urgent": cp.getint("output", "urgent_days", fallback=30),
        "preview": cp.getint("output", "preview_rows", fallback=7),
        "page_size": min(cp.getint("network", "page_size", fallback=100), 100),
        "delay": cp.getfloat("network", "delay", fallback=0.3),
        "retries": cp.getint("network", "retries", fallback=4),
        "batch_size": cp.getint("network", "batch_size", fallback=100),
        "batch_pause": cp.getfloat("network", "batch_pause", fallback=30),
        "max_requests": cp.getint("network", "max_requests", fallback=0),
        "cache_dir": cp.get("cache", "cache_dir", fallback="cache").strip(),
        "ttl_days": cp.getint("cache", "details_ttl_days", fallback=30),
        "dedupe": cp.getboolean("cache", "dedupe_applicants", fallback=True),
        "proxy": cp.get("network", "proxy", fallback="").strip(),
        "chunk_days": max(cp.getint("network", "chunk_days", fallback=3), 1),
    }
    if not cfg["kinds"]:
        sys.exit("В config.ini не указаны document_types (certificates / declarations).")
    return cfg


# ----------------------------------------------------------------------------
# Клиент API ФСА
# ----------------------------------------------------------------------------
class ApiBadRequest(RuntimeError):
    """Сайт ответил 400 - повторять тот же запрос бессмысленно."""


class NetworkGiveUp(RuntimeError):
    """Сеть/сайт не отвечает (вероятно, блокировка IP). Дальше стучаться нельзя."""


class RunLimit(NetworkGiveUp):
    """Исчерпан лимит запросов за один запуск (max_requests в config.ini)."""


class FsaClient:
    def __init__(self, cfg):
        self.delay = cfg["delay"]
        self.retries = max(cfg["retries"], 1)
        self.batch_size = cfg["batch_size"]
        self.batch_pause = cfg["batch_pause"]
        self.max_requests = cfg["max_requests"]
        self.made = 0
        self.s = requests.Session()
        proxy = (cfg["proxy"] or "").strip()
        if proxy.lower() == "none":              # принудительно без прокси
            self.s.trust_env = False
        elif proxy:                              # свой прокси из config.ini
            self.s.proxies = {"http": proxy, "https": proxy}
        self.s.headers.update({
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                          "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "Origin": BASE,
            "Referer": f"{BASE}/rss/certificate",
        })
        self.token = None

    def login(self):
        r = self.s.post(f"{BASE}/login", json=ANON_LOGIN, timeout=30)
        r.raise_for_status()
        token = r.headers.get("Authorization") or r.headers.get("authorization")
        if not token:
            raise NetworkGiveUp("Сайт не вернул токен авторизации (блокировка или изменился вход)")
        self.token = token
        self.s.headers["Authorization"] = token

    def _pace(self):
        """Неравномерные паузы + длинный перерыв после каждой пачки запросов."""
        self.made += 1
        time.sleep(self.delay * random.uniform(0.6, 1.6))
        if self.batch_size and self.made % self.batch_size == 0 and self.batch_pause > 0:
            sys.stdout.write(f"\r   {paint(f'Пауза {self.batch_pause:.0f} с, чтобы не нагружать сайт...', 'd')}"
                             + " " * 20)
            sys.stdout.flush()
            time.sleep(self.batch_pause)

    def _request(self, method, url, **kw):
        if self.max_requests and self.made >= self.max_requests:
            raise RunLimit(f"достигнут лимит {self.max_requests} запросов за запуск (max_requests)")
        last = None
        for attempt in range(1, self.retries + 1):
            wait = min(5 * 3 ** (attempt - 1), 120)       # 5, 15, 45, 120 с
            try:
                if not self.token:
                    self.login()
                r = self.s.request(method, url, timeout=60, **kw)
                sc = r.status_code
                if sc == 400:
                    raise ApiBadRequest(r.text[:300] or "HTTP 400")
                if sc in (401, 403):
                    self.token = None                      # токен истёк - войдём заново
                    last = RuntimeError(f"HTTP {sc}")
                elif sc == 429 or sc >= 500:
                    last = RuntimeError(f"HTTP {sc}")
                    ra = r.headers.get("Retry-After", "")
                    if ra.isdigit():
                        wait = min(int(ra), 300)
                else:
                    r.raise_for_status()
                    data = r.json()
                    self._pace()
                    return data
            except requests.exceptions.ProxyError as e:
                raise NetworkGiveUp(
                    "Не удаётся подключиться к прокси (он не запущен или указан неверно).\n"
                    "  Запустите прокси/VPN либо задайте параметр proxy в config.ini:\n"
                    "    proxy =                        - брать из системы (HTTPS_PROXY)\n"
                    "    proxy = none                   - работать без прокси\n"
                    "    proxy = http://127.0.0.1:10809 - указать прокси явно\n"
                    f"  Детали: {str(e)[:200]}")
            except (requests.RequestException, ValueError) as e:
                last = e
            if attempt < self.retries:
                time.sleep(wait)
        raise NetworkGiveUp(f"сайт не отвечает ({str(last)[:150]}). Возможно, IP временно заблокирован")

    def list_page(self, kind, page, size, d_from, d_to, tr_ts_id=None):
        body = {
            "size": size,
            "page": page,
            "filter": {
                "status": [STATUS_ACTIVE],
                "endDate": {"minDate": d_from.isoformat(), "maxDate": d_to.isoformat()},
                "regDate": {"minDate": None, "maxDate": None},
                "idCertType": [], "idObjectType": [], "idProductType": [],
                "idGroupRU": [], "idGroupEEU": [], "idTechReg": [tr_ts_id] if tr_ts_id else [], "idApplicantType": [],
                "idCertScheme": [], "idDeclScheme": [], "idProductOrigin": [],
                "idProductEEU": [], "idProductRU": [],
                "columnsSearch": [],
                "awaitForApprove": None, "awaitOperatorCheck": None, "editApp": None,
                "violationSend": None, "isProtected": None, "isNotAccredited": None,
            },
            "columnsSort": [{"column": DOC_KINDS[kind]["sort_column"], "sort": "ASC"}],
        }
        return self._request("POST", DOC_KINDS[kind]["list_url"], data=json.dumps(body))

    def detail(self, kind, doc_id):
        return self._request("GET", DOC_KINDS[kind]["detail_url"].format(id=doc_id))


# ----------------------------------------------------------------------------
# Разбор ответов (сделан устойчивым к небольшим отличиям в названиях полей)
# ----------------------------------------------------------------------------
def first_key(d, *names, suffix=None):
    """Первое непустое значение по списку имён ключей (или по окончанию ключа)."""
    if not isinstance(d, dict):
        return None
    for n in names:
        v = d.get(n)
        if v not in (None, "", []):
            return v
    if suffix:
        for k, v in d.items():
            if k.lower().endswith(suffix.lower()) and v not in (None, "", []):
                return v
    return None


def parse_date(v):
    if not v:
        return None
    if isinstance(v, (int, float)):                       # unix-время в мс
        return datetime.fromtimestamp(v / 1000).date()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(v))
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", str(v))
    if m:
        return date(int(m[3]), int(m[2]), int(m[1]))
    return None


def collect_strings(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for v in obj.values():
            yield from collect_strings(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from collect_strings(v)


CONTACT_EMAIL_TYPE = 4   # idContactType контакта «E-mail» на сайте


def find_email(applicant, whole):
    """Сначала контакты заявителя (тип e-mail), затем любой e-mail по карточке."""
    for src in (applicant, whole):
        if not isinstance(src, dict):
            continue
        for c in src.get("contacts") or []:         # точный путь: тип контакта 4 = e-mail
            if isinstance(c, dict) and c.get("idContactType") == CONTACT_EMAIL_TYPE:
                m = EMAIL_RE.search(str(c.get("value") or ""))
                if m:
                    return m.group(0).lower()
        for text in collect_strings(src.get("contacts") or []):
            m = EMAIL_RE.search(text)
            if m:
                return m.group(0).lower()
        for text in collect_strings(src):
            m = EMAIL_RE.search(text)
            if m:
                return m.group(0).lower()
    return ""


def find_head(applicant):
    """ФИО руководителя заявителя."""
    if not isinstance(applicant, dict):
        return ""
    direct = first_key(applicant, "headFio", "headName", "directorFio", "directorName",
                       "managerFio", "manager", "headFullName", "fioHead")
    if isinstance(direct, str):
        return direct.strip()
    if isinstance(direct, dict):
        applicant = direct
    parts = [applicant.get("surname"), applicant.get("firstName"), applicant.get("patronymic")]
    if any(parts):
        return " ".join(p.strip() for p in parts if p).strip()
    for k, v in applicant.items():                         # запасной вариант
        if isinstance(v, str) and any(w in k.lower() for w in ("head", "director", "fio")):
            return v.strip()
    return ""


def extract_applicant(det):
    """Из карточки документа берём только нужное: юрлицо, руководитель, e-mail, ТР ТС."""
    applicant = (det or {}).get("applicant") or {}
    return {
        "company": str(first_key(applicant, "fullName", "shortName", "name") or "").strip(),
        "head": find_head(applicant),
        "email": find_email(applicant, det),
        "tr_ts": extract_tr_ts(det),
    }


def applicant_key(item):
    """Ключ заявителя для общего кэша (ИНН, если есть в списке, иначе нормализованное имя)."""
    v = first_key(item, "applicantInn", "applicantINN", "inn", "idApplicant")
    if v:
        return f"id:{v}"
    n = " ".join(str(first_key(item, "applicantName") or "").lower().split())
    return f"n:{n}" if n else None


NAME_PARTICLES = {"оглы", "огли", "кызы", "кизи", "гызы", "угли", "улы", "кзы"}


def clean_name(name, mode="title"):
    """title - ЗАГЛАВНЫЕ буквы превращаем в обычный регистр; registry - как в реестре."""
    name = " ".join(str(name or "").split())
    if mode == "registry":
        return name
    out = []
    for w in name.split(" "):
        if w.lower() in NAME_PARTICLES:
            out.append(w.lower())
        elif w.isupper():
            out.append(w.title())
        else:
            out.append(w)
    return " ".join(out)


def split_name(name):
    """В реестре порядок: Фамилия Имя Отчество. 'Мамед оглы' остаётся в отчестве целиком."""
    parts = name.split()
    if not parts:
        return "", "", ""
    return parts[0], (parts[1] if len(parts) > 1 else ""), " ".join(parts[2:])


def build_row(kind, item, info, today, name_case="title", tr_ts="", tr_ts_doc=False):
    info = info or {}
    head = clean_name(info.get("head", ""), name_case)
    last_name, first_name, patronymic = split_name(head)
    end = parse_date(first_key(item, "certEndDate", "declEndDate", "endDate", suffix="EndDate"))
    company = info.get("company") or first_key(item, "applicantName") or ""
    product = first_key(item, "productFullName", "productName") or ""
    # ТР ТС: по умолчанию тот, по которому документ найден в реестре;
    # если в config.ini включено tr_ts_from_document - берём из карточки документа.
    tr_ts_value = ((info.get("tr_ts") or "") if tr_ts_doc else "") or tr_ts
    return {
        "company": str(company).strip(),
        "head": head,
        "last_name": last_name,
        "first_name": first_name,
        "patronymic": patronymic,
        "number": str(first_key(item, "number") or "").strip(),
        "doc_type": DOC_KINDS[kind]["title"],
        "end_date": end,
        "days_left": (end - today).days if end else None,
        "email": info.get("email", ""),
        "tr_ts": tr_ts_value,
        "product": str(product).strip().replace("\n", " "),
        "_has_info": bool(info),
    }


# ----------------------------------------------------------------------------
# Кэш на диске: найденное не пропадает при обрыве, повторный запуск продолжает с места остановки
# ----------------------------------------------------------------------------
class Cache:
    def __init__(self, folder, ttl_days):
        Path(folder).mkdir(parents=True, exist_ok=True)
        self.path = Path(folder) / "fsa_cache.json"
        self.data = {"windows": {}, "details": {}, "apps": {}}
        if self.path.exists():
            try:
                self.data.update(json.loads(self.path.read_text(encoding="utf-8")))
            except Exception:
                pass                                   # битый кэш просто игнорируем
        now, ttl, today = time.time(), ttl_days * 86400, date.today().isoformat()
        self.data["windows"] = {k: v for k, v in self.data["windows"].items() if k.endswith("|" + today)}
        for sect in ("details", "apps"):
            self.data[sect] = {k: v for k, v in self.data[sect].items() if now - v.get("t", 0) < ttl}
        self._dirty = 0

    def save(self):
        tmp = self.path.with_suffix(".tmp")
        tmp.write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8")
        os.replace(tmp, self.path)                     # запись атомарная - файл не испортится при обрыве
        self._dirty = 0

    def put_window(self, key, items):
        self.data["windows"][key] = items
        self.save()

    def get_info(self, doc_key, app_key):
        v = self.data["details"].get(doc_key) or (self.data["apps"].get(app_key) if app_key else None)
        return {k: v[k] for k in ("company", "head", "email", "tr_ts") if k in v} if v else None

    def put_info(self, doc_key, app_key, info):
        rec = dict(info, t=time.time())
        self.data["details"][doc_key] = rec
        if app_key:
            self.data["apps"][app_key] = rec
        self._dirty += 1
        if self._dirty >= 20:
            self.save()


# ----------------------------------------------------------------------------
# Сбор данных
# ----------------------------------------------------------------------------
def progress(done, total, label, unit=None):
    if unit:   # done - найдено документов, total - пройдено дней
        sys.stdout.write(f"\r   {paint(label, 'd')} найдено {done}, пройдено {total} {unit}      ")
        sys.stdout.flush()
        return
    total = max(total, 1)
    pct = int(done * 100 / total)
    sys.stdout.write(f"\r   {paint(label, 'd')} {done:>6}/{total:<6} ({pct:>3}%)   ")
    sys.stdout.flush()


MAX_WINDOW = 1500   # если в окне больше записей - делим окно пополам


def absorb(cfg, items, acc):
    for it in items:
        if it.get("id") in acc["seen"]:
            continue
        acc["seen"].add(it.get("id"))
        name = str(first_key(it, "productFullName", "productName") or "").lower()
        if not cfg["keywords"] or any(k in name for k in cfg["keywords"]):
            acc["matched"].append(it)


def fetch_window(client, kind, cfg, cache, a, b, acc, state, tr_ts="", tr_ts_id=None):
    """Читает все страницы за окно дат [a, b] по одному техрегламенту."""
    key = f"{kind}|{tr_ts}|{a}|{b}|{date.today()}"
    cached = cache.data["windows"].get(key)
    if cached is not None:                       # окно уже загружено ранее - сайт не трогаем
        absorb(cfg, cached, acc)
        return

    def split():
        span = (b - a).days
        if span >= 2:
            mid = a + timedelta(days=span // 2)
            parts = [(a, mid), (mid, b)]         # перекрытие на 1 день - дубли снимает seen
        else:
            parts = [(a, a), (b, b)]
        for x, y in parts:
            fetch_window(client, kind, cfg, cache, x, y, acc, state, tr_ts, tr_ts_id)

    page, got = 0, []
    while True:
        try:
            resp = client.list_page(kind, page, cfg["page_size"], a, b, tr_ts_id)
        except ApiBadRequest as e:
            if a < b:
                return split()
            state["warnings"].append(f"{kind} {a:%d.%m.%Y}, стр. {page}: {e}")
            return
        total = resp.get("total") or 0
        if page == 0 and total > MAX_WINDOW and a < b:
            return split()
        items = resp.get("items") or []
        got += items
        absorb(cfg, items, acc)
        progress(len(acc["seen"]), state["done_days"], f"{DOC_KINDS[kind]['title']} {a:%d.%m}:", unit="дн.")
        if not items or (page + 1) * cfg["page_size"] >= total:
            cache.put_window(key, got)
            return
        page += 1


def fetch_candidates(client, kind, cfg, cache, d_from, d_to, acc, state, tr_ts="", tr_ts_id=None):
    """Читает реестр окнами по нескольку дней. Результат накапливается в acc, даже если упадём на полпути."""
    cur = d_from
    while cur <= d_to:
        end = min(cur + timedelta(days=cfg["chunk_days"]), d_to)
        fetch_window(client, kind, cfg, cache, cur, end, acc, state, tr_ts, tr_ts_id)
        state["done_days"] = (end - d_from).days
        cur = end + timedelta(days=1)
    sys.stdout.write("\r" + " " * 78 + "\r")


def fmt_int(n):
    return f"{n:,}".replace(",", " ")


def run(cfg, debug=False):
    today = date.today()
    d_from, d_to = today, today + timedelta(days=cfg["days"])
    client = FsaClient(cfg)
    cache = Cache(cfg["cache_dir"], cfg["ttl_days"])
    state = {"warnings": [], "done_days": 0, "incomplete": False, "reason": "", "missing": 0}

    def stop(reason):
        if not state["incomplete"]:
            state["reason"] = reason
        state["incomplete"] = True
        sys.stdout.write("\r" + " " * 78 + "\r")

    print(paint("-" * 64, "d"))
    print(paint(f"Реестр ФСА: сертификаты и декларации о соответствии (версия {VERSION})", "d"))
    print()

    pool, infos = [], {}
    try:
        # 1-2. Загрузка реестра и фильтр по техрегламенту
        print(f" {paint('1.', 'b')} Загружаю реестр (истекают {d_from:%d.%m.%Y} - {d_to:%d.%m.%Y})")
        total_loaded = 0
        for tr in (cfg["tr_ts"] or [""]):           # пусто - весь реестр без фильтра по ТР ТС
            for kind in cfg["kinds"]:
                acc = {"seen": set(), "matched": []}
                try:
                    fetch_candidates(client, kind, cfg, cache, d_from, d_to, acc, state, tr, TR_TS_IDS.get(tr))
                finally:
                    total_loaded += len(acc["seen"])
                    pool += [(kind, it, tr) for it in acc["matched"]]
                    if debug and acc["matched"]:
                        tag = "_" + tr.replace("/", "-").replace(" ", "") if tr else ""
                        Path(f"debug_list_{kind}{tag}.json").write_text(
                            json.dumps(acc["matched"][0], ensure_ascii=False, indent=2), encoding="utf-8")
        for w in state["warnings"]:
            print(paint(f"    ВНИМАНИЕ, данные могли быть неполными: {w}", "y"))
        print(f"    ... {paint(fmt_int(total_loaded) + ' документов', 'd')}")
        flt = ", ".join(cfg["tr_ts"]) or "без фильтра по техрегламенту"
        print(f" {paint('2.', 'b')} Отбираю по техническому регламенту ({flt})")
        if cfg["keywords"]:
            print(f"    ... {paint('дополнительно по продукции: ' + ', '.join(cfg['keywords']), 'd')}")
        print(f"    ... {paint(fmt_int(len(pool)) + ' документов', 'd')}")
    except NetworkGiveUp as e:
        stop(str(e))
    except KeyboardInterrupt:
        stop("прервано пользователем (Ctrl+C)")
    except Exception as e:
        stop(f"непредвиденная ошибка: {e}")

    # 3. Карточки заявителей (руководитель, e-mail): сначала кэш, потом сайт
    print(f" {paint('3.', 'b')} Получаю данные заявителей (руководитель, e-mail)")
    online = not state["incomplete"]               # если на этапе 1 сломались - сайт больше не трогаем
    hits = fetched = 0
    try:
        for i, (kind, it, tr) in enumerate(pool, 1):
            dk, ak = f"{kind}:{it['id']}", (applicant_key(it) if cfg["dedupe"] else None)
            info = cache.get_info(dk, ak)
            if info:
                hits += 1
            elif online:
                try:
                    det = client.detail(kind, it["id"])
                except NetworkGiveUp:
                    raise
                except Exception as e:             # одна неудача не должна ронять всю выгрузку
                    print(paint(f"\n    не удалось получить {it.get('number')}: {e}", "r"))
                    continue
                info = extract_applicant(det)
                fetched += 1
                if debug and fetched == 1:
                    Path(f"debug_detail_{kind}.json").write_text(
                        json.dumps(det, ensure_ascii=False, indent=2), encoding="utf-8")
                if any(info.values()):             # пустой результат не кэшируем
                    cache.put_info(dk, ak, info)
            if info:
                infos[dk] = info
            progress(i, len(pool), "Заявители:")
    except NetworkGiveUp as e:
        stop(str(e))
    except KeyboardInterrupt:
        stop("прервано пользователем (Ctrl+C)")
    except Exception as e:
        stop(f"непредвиденная ошибка: {e}")
    finally:
        cache.save()
    sys.stdout.write("\r" + " " * 78 + "\r")
    print(f"    ... из кэша {hits}, запрошено с сайта {fetched}")

    rows = []
    for kind, it, tr in pool:
        info = infos.get(f"{kind}:{it['id']}")
        row = build_row(kind, it, info, today, cfg["name_case"], tr, cfg["tr_ts_doc"])
        if row["end_date"] and 0 <= row["days_left"] <= cfg["days"]:
            rows.append(row)
    rows.sort(key=lambda r: (r["days_left"], r["company"]))
    state["missing"] = sum(1 for r in rows if not r["_has_info"])
    print(f"    ... {paint(fmt_int(len(rows)) + ' документов', 'y')}")
    return rows, state


# ----------------------------------------------------------------------------
# Вывод в файлы
# ----------------------------------------------------------------------------
def fmt_value(col, v):
    if col == "end_date":
        return v.strftime("%d.%m.%Y") if v else ""
    return "" if v is None else v


def unique_path(folder, stem, ext):
    p = folder / f"{stem}.{ext}"
    n = 2
    while p.exists():                              # старые выгрузки не перезаписываем
        p = folder / f"{stem}_{n}.{ext}"
        n += 1
    return p


def save_csv(rows, cols, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:   # BOM - чтобы Excel не ломал кириллицу
        w = csv.writer(f, delimiter=";")
        w.writerow([COLUMN_TITLES[c] for c in cols])
        for r in rows:
            w.writerow([fmt_value(c, r[c]) for c in cols])


def save_xlsx(rows, cols, path, cfg):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Документы"
    font = Font(name="Arial", size=10)
    head_font = Font(name="Arial", size=10, bold=True)
    thin = Side(style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    red_fill = PatternFill("solid", fgColor="FDE8E8")

    ws.append([COLUMN_TITLES[c] for c in cols])
    for j in range(1, len(cols) + 1):
        c = ws.cell(row=1, column=j)
        c.font, c.border = head_font, border
        c.fill = PatternFill("solid", fgColor="EDEFF3")
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for r in rows:
        ws.append([r[c] if c in ("end_date", "days_left") else fmt_value(c, r[c]) for c in cols])
        i = ws.max_row
        urgent = r["days_left"] is not None and r["days_left"] <= cfg["urgent"]
        for j, col in enumerate(cols, 1):
            cell = ws.cell(row=i, column=j)
            cell.font, cell.border = font, border
            cell.alignment = Alignment(vertical="top", wrap_text=(col in ("company", "product", "tr_ts")))
            if col == "end_date":
                cell.number_format = "DD.MM.YYYY"
            if col in ("end_date", "days_left"):
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if urgent:
                cell.fill = red_fill
                if col == "days_left":
                    cell.font = Font(name="Arial", size=10, bold=True, color="C00000")

    for j, col in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = COLUMN_WIDTHS[col]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 24
    wb.save(path)


def save(rows, cfg, partial=False):
    cols = cfg["columns"]
    folder = Path(cfg["out_dir"])
    folder.mkdir(parents=True, exist_ok=True)
    stem = f"{cfg['prefix']}_{cfg['days']}_days_{date.today():%Y-%m-%d}" + ("_частично" if partial else "")
    paths = []
    if cfg["format"] in ("xlsx", "both"):
        p = unique_path(folder, stem, "xlsx")
        save_xlsx(rows, cols, p, cfg)
        paths.append(p)
    if cfg["format"] in ("csv", "both"):
        p = unique_path(folder, stem, "csv")
        save_csv(rows, cols, p)
        paths.append(p)
    return paths


def preview(rows, cfg):
    n = min(cfg["preview"], len(rows))
    if not n:
        return
    print(f"\n Первые {n} строк для контроля:")
    cols = [c for c in cfg["columns"] if c in ("company", "number", "tr_ts", "end_date", "days_left", "email")]
    widths = {"company": 28, "number": 30, "tr_ts": 16, "end_date": 11, "days_left": 6, "email": 28}
    short = {"company": "Юрлицо", "number": "Номер", "tr_ts": "ТР ТС", "end_date": "Окончание",
             "days_left": "Дней", "email": "E-mail"}
    print(paint("  " + " ".join(short[c].ljust(widths[c]) for c in cols), "d"))
    for r in rows[:n]:
        line = " ".join(str(fmt_value(c, r[c]))[:widths[c]].ljust(widths[c]) for c in cols)
        print("  " + paint(line, "r") if r["days_left"] <= cfg["urgent"] else "  " + line)


# ----------------------------------------------------------------------------
def print_tr_ts():
    """Список техрегламентов, доступных для фильтра."""
    print("Технические регламенты для параметра tr_ts в config.ini:")
    print()
    for label, tid in TR_TS_IDS.items():
        print(f"  {label:<42} id {tid}")
    print()
    print("Можно указывать без года: 017 (найдёт 017/2011).")
    print("Несколько - через запятую: tr_ts = 017, 021")
    print("Пустое значение - без фильтра по техрегламенту.")


def main():
    ap = argparse.ArgumentParser(description="Парсер реестра ФСА")
    ap.add_argument("--config", default=str(Path(__file__).with_name("config.ini")))
    ap.add_argument("--days", type=int, help="порог в днях (перекрывает config.ini)")
    ap.add_argument("--tr-ts", dest="tr_ts",
                    help="техрегламент: 017 или 017,021; none - без фильтра (перекрывает config.ini)")
    ap.add_argument("--list-trts", action="store_true", help="показать доступные техрегламенты и выйти")
    ap.add_argument("--debug", action="store_true", help="сохранить сырые ответы сайта")
    args = ap.parse_args()

    if args.list_trts:
        print_tr_ts()
        return

    cfg = load_config(args.config)
    if args.days:
        cfg["days"] = args.days
    if args.tr_ts:
        if args.tr_ts.strip().lower() in ("none", "нет", "все", "all", "-"):
            cfg["tr_ts"] = []
        else:
            labels = []
            for raw in args.tr_ts.replace(";", ",").split(","):
                if not raw.strip():
                    continue
                label = normalize_tr_ts(raw)
                if not label:
                    sys.exit(f"Не понял технический регламент «{raw}».\n"
                             f"  Список: python fsa_parser.py --list-trts")
                if label not in labels:
                    labels.append(label)
            cfg["tr_ts"] = labels

    rows, state = run(cfg, debug=args.debug)
    partial = state["incomplete"] or state["missing"] > 0

    if not rows:
        print(paint("\nДокументов для выгрузки нет - файл не создан.", "y"))
    else:
        print(f" {paint('4.', 'b')} Сохраняю файл")
        paths = save(rows, cfg, partial)
        print(f"    ... {paint('готово', 'g')}")
        preview(rows, cfg)
        print()
        label = "Файл ЧАСТИЧНЫЙ" if partial else "Файл готов"
        print(paint(f"Найдено {fmt_int(len(rows))} документов. {label}.", "y" if partial else "g"))
        if cfg["tr_ts"]:
            print(paint("ТР ТС: " + ", ".join(cfg["tr_ts"]), "d"))
        for p in paths:
            print(paint(str(p.resolve()), "d"))

    if partial:
        print()
        if state["incomplete"]:
            print(paint(f"Выгрузка остановлена: {state['reason']}", "r"))
        if state["missing"]:
            print(paint(f"У {fmt_int(state['missing'])} документов нет руководителя/e-mail (не успели загрузить).", "y"))
        print("Всё найденное сохранено в кэше. Подождите (при блокировке - несколько часов),")
        print("запустите парсер снова - он продолжит с места остановки, а не начнёт заново.")
        sys.exit(2)


if __name__ == "__main__":
    main()